import csv
import io
import re
from io import BytesIO

import openpyxl
import pdfplumber
from docx import Document
from sqlalchemy.orm import Session

from app.models import ExamCenter, Student, Upload

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")

# Substrings matched against a header cell after punctuation is stripped, so
# "Phone No.", "Mobile No", "Contact Number", "Email ID", "Candidate Name",
# etc. are all recognized without needing an exact alias match.
FIELD_KEYWORDS = {
    "name": ["name"],
    "email": ["email", "e mail"],
    "mobile": ["mobile", "phone", "contact", "whatsapp"],
}


def _normalize_header(cell) -> str:
    text = str(cell or "").strip().lower()
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _match_header(header_cells: list[str]) -> dict[str, int]:
    normalized = [_normalize_header(h) for h in header_cells]
    col_map: dict[str, int] = {}
    for field, keywords in FIELD_KEYWORDS.items():
        for idx, cell in enumerate(normalized):
            if idx in col_map.values():
                continue
            if any(keyword in cell for keyword in keywords):
                col_map[field] = idx
                break
    missing = [f for f in FIELD_KEYWORDS if f not in col_map]
    if missing:
        raise ValueError(
            f"Sheet is missing required column(s): {', '.join(missing)}. "
            f"Expected headers like: Student Name, Email, Mobile Number/Phone No."
        )
    return col_map


def _cell(row: tuple, idx: int) -> str:
    if idx >= len(row):
        return ""
    return str(row[idx] if row[idx] is not None else "").strip()


def _is_blank_row(row: tuple) -> bool:
    return not any(str(c or "").strip() for c in row)


def _decode_text(file_bytes: bytes) -> str:
    """CSV files exported from Excel/Google Sheets on Windows are frequently
    saved as cp1252/latin-1 rather than UTF-8, and often carry a BOM — try
    the common encodings in order rather than assuming UTF-8 and failing
    the whole upload over a single rupee/mojibake character."""
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="replace")


def _rows_from_csv(file_bytes: bytes) -> list[tuple]:
    """Parses a .csv upload. Delimiter is sniffed rather than assumed —
    files saved from Excel as "CSV" sometimes actually use tabs or
    semicolons — and fully blank rows are dropped up front."""
    text = _decode_text(file_bytes)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows = [tuple(row) for row in reader if not _is_blank_row(row)]
    if not rows:
        raise ValueError("CSV file is empty.")
    return rows


def _rows_from_xlsx(file_bytes: bytes) -> list[tuple]:
    wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
    ws = wb.active
    all_rows = list(ws.iter_rows(values_only=True))
    if not all_rows:
        raise ValueError("Excel sheet is empty.")
    header, *data_rows = all_rows
    rows = [header, *(r for r in data_rows if not _is_blank_row(r))]
    return rows


def _rows_from_pdf(file_bytes: bytes) -> list[tuple]:
    """Extracts a student table from a PDF. Assumes the same header row as the
    Excel path (Name/Email/Mobile); if that header repeats on every page (common
    when a table spans multiple pages), later occurrences are treated as a
    repeated header and skipped rather than as a data row."""
    rows: list[tuple] = []
    header_key: list[str] | None = None

    with pdfplumber.open(BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables():
                if not table:
                    continue
                table_header, *table_rows = table
                normalized = [str(c or "").strip().lower() for c in table_header]

                if header_key is None:
                    header_key = normalized
                    rows.append(tuple(table_header))
                    rows.extend(tuple(r) for r in table_rows if not _is_blank_row(r))
                elif normalized == header_key:
                    rows.extend(tuple(r) for r in table_rows if not _is_blank_row(r))
                else:
                    rows.extend(tuple(r) for r in [table_header, *table_rows] if not _is_blank_row(r))

    if not rows:
        raise ValueError("Could not find a table in the PDF.")
    return rows


def _rows_from_docx(file_bytes: bytes) -> list[tuple]:
    """Extracts a student table from a Word document, same header-repeat
    handling as the PDF path in case the document has more than one table
    (e.g. a table per page)."""
    rows: list[tuple] = []
    header_key: list[str] | None = None

    doc = Document(BytesIO(file_bytes))
    for table in doc.tables:
        if not table.rows:
            continue
        table_rows = [[cell.text for cell in row.cells] for row in table.rows]
        table_header, *data_rows = table_rows
        normalized = [str(c or "").strip().lower() for c in table_header]

        if header_key is None:
            header_key = normalized
            rows.append(tuple(table_header))
            rows.extend(tuple(r) for r in data_rows if not _is_blank_row(r))
        elif normalized == header_key:
            rows.extend(tuple(r) for r in data_rows if not _is_blank_row(r))
        else:
            rows.extend(tuple(r) for r in [table_header, *data_rows] if not _is_blank_row(r))

    if not rows:
        raise ValueError("Could not find a table in the Word document.")
    return rows


def import_students_excel(
    db: Session, file_bytes: bytes, filename: str, uploaded_by_user_id: int, exam_center_id: int
) -> tuple[Upload, list[str]]:
    center = db.get(ExamCenter, exam_center_id)
    if center is None:
        raise ValueError("Unknown exam center.")

    lower_name = filename.lower()
    if lower_name.endswith(".csv"):
        rows = _rows_from_csv(file_bytes)
    elif lower_name.endswith(".pdf"):
        rows = _rows_from_pdf(file_bytes)
    elif lower_name.endswith(".docx"):
        rows = _rows_from_docx(file_bytes)
    else:
        rows = _rows_from_xlsx(file_bytes)

    col_map = _match_header(list(rows[0]))
    data_rows = [r for r in rows[1:] if not _is_blank_row(r)]

    new_count = 0
    updated_count = 0
    errors: list[str] = []

    for i, row in enumerate(data_rows, start=2):
        name = _cell(row, col_map["name"])
        email = _cell(row, col_map["email"]).lower()
        mobile = _cell(row, col_map["mobile"])

        if not name or not email or not mobile:
            errors.append(f"Row {i}: missing required value(s), skipped.")
            continue
        if not EMAIL_RE.match(email):
            errors.append(f"Row {i}: invalid email '{email}', skipped.")
            continue

        # Each row gets its own SAVEPOINT: a duplicate/invalid row that fails
        # at flush (e.g. two rows in the same sheet sharing an email) would
        # otherwise abort the whole Postgres transaction and silently lose
        # every row after it. Rolling back just this row's savepoint keeps
        # the rest of the batch — and the final commit — intact.
        try:
            with db.begin_nested():
                existing = db.query(Student).filter(Student.email == email).one_or_none()
                if existing:
                    existing.full_name = name
                    existing.mobile_number = mobile
                    existing.exam_center_id = center.id
                    updated_count += 1
                else:
                    db.add(
                        Student(
                            full_name=name,
                            email=email,
                            mobile_number=mobile,
                            exam_center_id=center.id,
                        )
                    )
                    new_count += 1
        except Exception as exc:  # noqa: BLE001 - keep importing remaining rows
            errors.append(f"Row {i}: {exc}")

    upload = Upload(
        filename=filename,
        uploaded_by=uploaded_by_user_id,
        row_count=len(data_rows),
        new_students=new_count,
        updated_students=updated_count,
        error_count=len(errors),
    )
    db.add(upload)
    db.commit()
    db.refresh(upload)
    return upload, errors
